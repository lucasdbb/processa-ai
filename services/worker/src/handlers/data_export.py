from __future__ import annotations

import hashlib
import json
import uuid
from datetime import UTC, date, datetime, time
from decimal import Decimal
from io import BytesIO
from typing import Any

import httpx
from openpyxl import Workbook
from psycopg.rows import dict_row

from src.core.config import settings
from src.core.db import get_pool
from src.core.job_store import (
    add_job_event,
    create_document,
    mark_job_completed,
    mark_job_failed,
    mark_job_processing,
)
from src.core.logging import logger
from src.core.storage import StorageService

SENDGRID_API_URL = "https://api.sendgrid.com/v3/mail/send"

# Keep sheet names short (Excel limit: 31 chars) and stable.
TABLE_EXPORTS: list[tuple[str, str]] = [
    ("Perfil", "SELECT to_jsonb(p) AS row_data FROM public.profiles p WHERE p.id = %s"),
    (
        "Assinaturas",
        "SELECT to_jsonb(s) AS row_data FROM public.subscriptions s WHERE s.user_id::text = %s",
    ),
    (
        "Clientes",
        "SELECT to_jsonb(c) AS row_data FROM office_api.clients c WHERE c.user_id::text = %s",
    ),
    (
        "Processos",
        "SELECT to_jsonb(p) AS row_data FROM office_api.processes p WHERE p.user_id::text = %s",
    ),
    (
        "Servicos",
        "SELECT to_jsonb(s) AS row_data FROM office_api.services s WHERE s.user_id::text = %s",
    ),
    (
        "Honorarios",
        "SELECT to_jsonb(f) AS row_data FROM office_api.fees f WHERE f.user_id::text = %s",
    ),
    (
        "Pagamentos",
        "SELECT to_jsonb(fp) AS row_data FROM office_api.fee_payments fp WHERE fp.user_id::text = %s",
    ),
    ("Notas", "SELECT to_jsonb(n) AS row_data FROM office_api.notes n WHERE n.user_id::text = %s"),
    (
        "Timeline",
        "SELECT to_jsonb(t) AS row_data FROM office_api.timeline_events t WHERE t.user_id::text = %s",
    ),
    (
        "Calendario",
        "SELECT to_jsonb(e) AS row_data FROM office_api.events e WHERE e.user_id::text = %s",
    ),
    (
        "DocsEscritorio",
        "SELECT to_jsonb(d) AS row_data FROM office_api.documents d WHERE d.user_id::text = %s",
    ),
    (
        "DocumentosGerados",
        "SELECT to_jsonb(gd) AS row_data FROM public.generated_documents gd WHERE gd.user_id::text = %s",
    ),
    (
        "Notificacoes",
        "SELECT to_jsonb(n) AS row_data FROM public.notifications n WHERE n.user_id::text = %s",
    ),
    (
        "Calculos",
        "SELECT to_jsonb(c) AS row_data FROM public.calculations c WHERE c.user_id::text = %s",
    ),
    ("Jobs", "SELECT to_jsonb(j) AS row_data FROM public.jobs j WHERE j.user_id::text = %s"),
]


def _to_jsonb(value: Any) -> Any:
    try:
        from psycopg.types.json import Jsonb
    except ModuleNotFoundError:
        return value
    return Jsonb(value)


def _safe_cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        text = value if isinstance(value, str) else value
        if isinstance(text, str) and len(text) > 32000:
            return text[:32000] + "...[truncated]"
        return text
    if isinstance(value, (datetime, date, time, uuid.UUID, Decimal)):
        return str(value)
    if isinstance(value, (dict, list, tuple)):
        text = json.dumps(value, ensure_ascii=False, default=str)
        if len(text) > 32000:
            return text[:32000] + "...[truncated]"
        return text
    text = str(value)
    if len(text) > 32000:
        return text[:32000] + "...[truncated]"
    return text


def _append_dataset_sheet(wb: Workbook, sheet_name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=sheet_name[:31])
    if not rows:
        ws.append(["info"])
        ws.append(["Sem registros"])
        return

    headers: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.add(key)
                headers.append(key)

    ws.append(headers)
    for row in rows:
        ws.append([_safe_cell_value(row.get(header)) for header in headers])

    ws.freeze_panes = "A2"


def _build_workbook(
    user_id: str,
    datasets: list[tuple[str, list[dict[str, Any]]]],
    errors: list[str],
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    meta_sheet = wb.create_sheet(title="Meta")
    meta_sheet.append(["field", "value"])
    meta_sheet.append(["user_id", user_id])
    meta_sheet.append(["generated_at", datetime.now(UTC).isoformat()])
    meta_sheet.append(["format", "xlsx"])

    if errors:
        errors_sheet = wb.create_sheet(title="Erros")
        errors_sheet.append(["erro"])
        for err in errors:
            errors_sheet.append([err])

    for sheet_name, rows in datasets:
        _append_dataset_sheet(wb, sheet_name, rows)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


async def _fetch_profile_data(user_id: str) -> dict[str, Any] | None:
    pool = await get_pool()
    async with pool.connection() as conn:
        async with conn.cursor(row_factory=dict_row) as cur:
            await cur.execute(
                "SELECT id, email, name FROM public.profiles WHERE id = %s LIMIT 1",
                (user_id,),
            )
            return await cur.fetchone()


async def _collect_datasets(
    user_id: str,
) -> tuple[list[tuple[str, list[dict[str, Any]]]], list[str]]:
    datasets: list[tuple[str, list[dict[str, Any]]]] = []
    errors: list[str] = []

    pool = await get_pool()
    async with pool.connection() as conn:
        async with conn.cursor(row_factory=dict_row) as cur:
            for sheet_name, query in TABLE_EXPORTS:
                try:
                    await cur.execute(query, (user_id,))
                    rows = await cur.fetchall()
                    normalized_rows: list[dict[str, Any]] = []
                    for row in rows:
                        data = row.get("row_data") if isinstance(row, dict) else None
                        if isinstance(data, dict):
                            normalized_rows.append(data)
                    datasets.append((sheet_name, normalized_rows))
                except Exception as exc:
                    logger.exception(
                        "Failed collecting dataset %s for user %s", sheet_name, user_id
                    )
                    datasets.append((sheet_name, []))
                    errors.append(f"{sheet_name}: {exc}")

    return datasets, errors


async def _create_export_notification(
    user_id: str,
    *,
    document_id: str | None,
    job_id: str,
    filename: str,
) -> None:
    pool = await get_pool()
    metadata: dict[str, Any] = {
        "job_id": job_id,
        "filename": filename,
    }
    if document_id:
        metadata["document_id"] = document_id

    async with pool.connection() as conn:
        async with conn.cursor() as cur:
            await cur.execute(
                """
                INSERT INTO public.notifications (user_id, type, title, description, metadata)
                VALUES (%s, %s, %s, %s, %s)
                """,
                (
                    user_id,
                    "system",
                    "Exportacao de dados pronta",
                    "Seu arquivo Excel esta pronto para download.",
                    _to_jsonb(metadata),
                ),
            )


async def _send_export_ready_email(
    *,
    to_email: str,
    user_name: str | None,
    filename: str,
    download_url: str | None,
) -> bool:
    api_key = (settings.SENDGRID_API_KEY or "").strip()
    sender = (settings.SENDGRID_FROM or "").strip()
    frontend_url = (settings.FRONTEND_URL or "http://localhost:5173").rstrip("/")

    if not api_key or not sender or not to_email:
        return False

    greeting_name = user_name.strip() if isinstance(user_name, str) and user_name.strip() else ""
    greeting = f"Ola, {greeting_name}!" if greeting_name else "Ola!"
    app_link = f"{frontend_url}/app/documentos"

    plain_body = (
        f"{greeting}\n\n"
        "Sua exportacao de dados foi concluida com sucesso.\n"
        f"Arquivo: {filename}\n"
        + (f"Download direto: {download_url}\n" if download_url else "")
        + f"Acesse tambem em: {app_link}\n\n"
        "Equipe Processa AI"
    )

    html_body = (
        f"<p>{greeting}</p>"
        "<p>Sua exportacao de dados foi concluida com sucesso.</p>"
        f"<p><strong>Arquivo:</strong> {filename}</p>"
        + (
            f'<p><a href="{download_url}">Clique aqui para baixar agora</a></p>'
            if download_url
            else ""
        )
        + f'<p>Voce tambem pode acessar em <a href="{app_link}">Documentos</a>.</p>'
        "<p>Equipe Processa AI</p>"
    )

    payload = {
        "personalizations": [
            {"to": [{"email": to_email}], "subject": "Seu arquivo de exportacao esta pronto"}
        ],
        "from": {"email": sender},
        "content": [
            {"type": "text/plain", "value": plain_body},
            {"type": "text/html", "value": html_body},
        ],
    }

    try:
        async with httpx.AsyncClient(timeout=15) as client:
            resp = await client.post(
                SENDGRID_API_URL,
                json=payload,
                headers={
                    "Authorization": f"Bearer {api_key}",
                    "Content-Type": "application/json",
                },
            )
            return resp.status_code in (200, 201, 202)
    except Exception:
        logger.exception("Failed sending export-ready email")
        return False


async def handle_user_data_export_job(job_row: dict) -> None:
    """Generate an XLSX export with all user data and notify by email when ready."""
    raw_job_id = job_row.get("id")
    payload = (
        job_row.get("request_payload") if isinstance(job_row.get("request_payload"), dict) else {}
    )

    try:
        job_uuid = raw_job_id if isinstance(raw_job_id, uuid.UUID) else uuid.UUID(str(raw_job_id))
    except Exception:
        logger.error("Invalid job id for user_data_export: %s", raw_job_id)
        return

    user_id = str(payload.get("user_id") or "").strip()
    if not user_id:
        await mark_job_failed(job_uuid=job_uuid, error="user_id obrigatorio")
        await add_job_event(job_uuid=job_uuid, event_type="failed", message="user_id obrigatorio")
        return

    try:
        await mark_job_processing(job_uuid=job_uuid, progress=10)
        await add_job_event(
            job_uuid=job_uuid,
            event_type="processing",
            message="Iniciando exportacao de dados",
            payload={"progress": 10},
        )

        profile = await _fetch_profile_data(user_id)
        user_email = str(
            payload.get("user_email") or (profile.get("email") if profile else "") or ""
        ).strip()
        user_name = str(
            payload.get("user_name") or (profile.get("name") if profile else "") or ""
        ).strip()

        datasets, errors = await _collect_datasets(user_id)

        await add_job_event(
            job_uuid=job_uuid,
            event_type="processing",
            message="Gerando arquivo Excel",
            payload={"progress": 55},
        )

        xlsx_bytes = _build_workbook(user_id, datasets, errors)
        filename = f"meus_dados_{datetime.now(UTC).strftime('%Y%m%d_%H%M%S')}.xlsx"

        storage = StorageService()
        uploaded = await storage.upload_bytes_to_supabase(
            file_bytes=xlsx_bytes,
            user_id=user_id,
            client_id=None,
            functionality="data_exports",
            job_uuid=str(job_uuid),
            filename=filename,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            expires_in=86400,
            output=True,
        )

        bucket = uploaded.get("bucket")
        storage_path = uploaded.get("storage_path")
        signed_url = uploaded.get("signed_url")

        document_id = None
        if isinstance(bucket, str) and isinstance(storage_path, str):
            checksum = hashlib.sha256(xlsx_bytes).hexdigest()

            document_id = await create_document(
                job_uuid=job_uuid,
                user_id=user_id,
                client_id=None,
                document_type="user_data_export_xlsx",
                bucket=bucket,
                storage_path=storage_path,
                filename=filename,
                checksum_sha256=checksum,
                size_bytes=len(xlsx_bytes),
                mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                metadata={"job_type": "user_data_export", "errors": errors},
                source="system",
            )

        result_payload: dict[str, Any] = {
            "filename": filename,
            "download_url": signed_url,
            "errors_count": len(errors),
        }
        if document_id:
            result_payload["document_id"] = str(document_id)

        await mark_job_completed(job_uuid=job_uuid, result_payload=result_payload)
        await add_job_event(
            job_uuid=job_uuid,
            event_type="completed",
            message="Exportacao concluida",
            payload={"progress": 100},
        )

        await _create_export_notification(
            user_id,
            document_id=str(document_id) if document_id else None,
            job_id=str(job_uuid),
            filename=filename,
        )

        if user_email:
            email_sent = await _send_export_ready_email(
                to_email=user_email,
                user_name=user_name or None,
                filename=filename,
                download_url=signed_url if isinstance(signed_url, str) else None,
            )
            if not email_sent:
                logger.warning("Could not send export-ready email to user %s", user_id)

    except Exception as exc:
        logger.exception("user_data_export failed for user %s", user_id)
        try:
            await mark_job_failed(job_uuid=job_uuid, error=str(exc))
            await add_job_event(
                job_uuid=job_uuid,
                event_type="failed",
                message=str(exc),
            )
        except Exception:
            logger.exception("Failed to persist user_data_export failure state")
